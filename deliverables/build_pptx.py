#!/usr/bin/env python3
# build_pptx.py
# Usage: python deliverables/build_pptx.py

import os
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    import cairosvg
except Exception as e:
    print('Missing dependency:', e)
    print('Run: pip install -r deliverables/requirements.txt')
    raise

REPO_ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = REPO_ROOT / 'deliverables'
CHARTS_DIR = DELIVERABLES / 'charts'

# Ensure output paths
OUTPUT_PPTX = DELIVERABLES / 'front_half_draft_v1.pptx'

# Convert SVG charts to PNG (if SVGs exist)
svg_png_pairs = []
for svg_name in ['global_manufacturing.svg','robot_density.svg','china_ai_manufacturing.svg']:
    svg_path = CHARTS_DIR / svg_name
    png_path = CHARTS_DIR / svg_name.replace('.svg','.png')
    if svg_path.exists():
        try:
            cairosvg.svg2png(url=str(svg_path), write_to=str(png_path))
            print(f'Converted {svg_path.name} -> {png_path.name}')
            svg_png_pairs.append((svg_path.name, png_path.name))
        except Exception as e:
            print(f'Failed to convert {svg_path.name}:', e)
    else:
        print(f'SVG not found: {svg_path}')

# Read speaker notes file
speaker_file = DELIVERABLES / 'speaker_notes_part1.txt'
if not speaker_file.exists():
    raise FileNotFoundError(f'{speaker_file} not found. Ensure deliverables/speaker_notes_part1.txt exists')

text = speaker_file.read_text(encoding='utf-8')
# Split into sections by '第{n}页'
sections = re.split(r'(?=第\d+页)', text)
# Create presentation
prs = Presentation()
# Use Title and Content layout for slides
for sec in sections:
    sec = sec.strip()
    if not sec:
        continue
    lines = sec.splitlines()
    title = lines[0].strip()
    body = '\n'.join([ln.strip() for ln in lines[1:] if ln.strip()])
    slide_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    # set title if placeholder exists
    if slide.shapes.title:
        slide.shapes.title.text = title
    # set body text
    try:
        txBox = slide.shapes.placeholders[1].text_frame
        txBox.clear()
        for i, para in enumerate(body.split('\n\n')):
            p = txBox.add_paragraph() if i>0 else txBox.paragraphs[0]
            p.text = para.replace('\n',' ')
            p.font.size = Pt(12)
    except Exception:
        pass
    # add notes
    notes_slide = slide.notes_slide
    notes = notes_slide.notes_text_frame
    notes.text = body

# Appendix slide with charts
append_slide_layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[-1]
append_slide = prs.slides.add_slide(append_slide_layout)
append_title = '附录：核心图表'
if append_slide.shapes.title:
    append_slide.shapes.title.text = append_title
left = Inches(0.5)
top = Inches(1.2)
pic_w = Inches(4)
pic_h = Inches(3)
for i, (_, png_name) in enumerate(svg_png_pairs):
    png_path = CHARTS_DIR / png_name
    if png_path.exists():
        l = left + (i % 2) * (pic_w + Inches(0.5))
        t = top + (i // 2) * (pic_h + Inches(0.5))
        try:
            append_slide.shapes.add_picture(str(png_path), l, t, width=pic_w, height=pic_h)
        except Exception as e:
            print('Failed to add picture', png_path, e)

# Save presentation
prs.save(str(OUTPUT_PPTX))
print('Generated PPTX:', OUTPUT_PPTX)
